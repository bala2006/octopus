#!/usr/bin/env python3
"""Generate the Octopus pitch deck for Vivek (16:9, dark brand theme).

Usage:
    python build_pptx.py [output_path]

Produces: artifacts/octopus-pitch-for-vivek.pptx
Speaker notes are included on every slide.
"""

from __future__ import annotations

import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from content import PALETTE  # noqa: E402

# ---------------------------------------------------------------------------
# Layout constants
# ---------------------------------------------------------------------------
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
ML = Inches(0.8)                      # left margin
CW = Inches(13.333 - 1.6)             # content width
BODY_FONT = "Arial"                   # universal: renders identically everywhere
HEAD_FONT = "Arial"

TOTAL_SLIDES = 12


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#"))


# ---------------------------------------------------------------------------
# Primitives
# ---------------------------------------------------------------------------
def _no_shadow(shape) -> None:
    try:
        shape.shadow.inherit = False
    except Exception:  # pragma: no cover - defensive
        pass


def add_background(slide, deep: bool = False) -> None:
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shp.line.fill.background()
    _no_shadow(shp)
    fill = shp.fill
    try:
        fill.gradient()
        stops = fill.gradient_stops
        n = len(stops)
        stops[0].color.rgb = rgb(PALETTE["bg_deep"] if deep else PALETTE["bg"])
        stops[0].position = 0.0
        stops[n - 1].color.rgb = rgb("#111110")
        stops[n - 1].position = 1.0
        for i in range(1, n - 1):
            stops[i].color.rgb = rgb(PALETTE["bg"])
            stops[i].position = i / (n - 1)
        fill.gradient_angle = 45.0
    except Exception:
        fill.solid()
        fill.fore_color.rgb = rgb(PALETTE["bg_deep"] if deep else PALETTE["bg"])
    # faint decorative disc, bottom-right
    disc = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(10.7), Inches(4.9), Inches(3.6), Inches(3.6)
    )
    disc.fill.solid()
    disc.fill.fore_color.rgb = rgb("#26221F")
    disc.line.fill.background()
    _no_shadow(disc)


def add_rect(slide, x, y, w, h, color: str, rounded: bool = False, radius: float = 0.5):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, x, y, w, h)
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(color)
    shp.line.fill.background()
    _no_shadow(shp)
    return shp


def add_textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf


def add_run(p, text, size, color, *, bold=False, italic=False, font=BODY_FONT, spc=None):
    run = p.add_run()
    run.text = text
    f = run.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    if spc:
        try:
            run._r.get_or_add_rPr().set("spc", str(int(spc)))
        except Exception:
            pass
    return run


def add_line_of_text(slide, x, y, w, h, text, size, color, **kw):
    tf = add_textbox(slide, x, y, w, h)
    p = tf.paragraphs[0]
    p.line_spacing = kw.pop("line_spacing", 1.1)
    p.alignment = kw.pop("align", PP_ALIGN.LEFT)
    add_run(p, text, size, color, **kw)
    return tf


def add_kicker(slide, text: str) -> None:
    tf = add_textbox(slide, ML, Inches(0.62), CW, Inches(0.3))
    p = tf.paragraphs[0]
    add_run(p, text.upper(), 11, PALETTE["accent"], bold=True, spc=220)


def add_title(slide, text: str, size: float = 33) -> float:
    lines = text.split("\n")
    tf = add_textbox(slide, ML, Inches(0.95), CW, Inches(1.0 * len(lines)))
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.05
        add_run(p, line, size, PALETTE["text"], bold=True, font=HEAD_FONT, spc=-15)
    rule_y = Inches(0.95 + 0.62 * len(lines) + 0.16)
    add_rect(slide, ML, rule_y, Inches(1.15), Inches(0.055), PALETTE["accent"], rounded=True, radius=0.5)
    return float(rule_y / 914400) + 0.14  # next free y in inches


def add_footer(slide, page: int) -> None:
    add_rect(slide, ML, Inches(6.93), CW, Inches(0.011), "#37342F")
    tf = add_textbox(slide, ML, Inches(7.03), Inches(7.0), Inches(0.3))
    add_run(tf.paragraphs[0], "OCTOPUS  \u00b7  PITCH FOR VIVEK", 8.5, PALETTE["dim"], spc=140)
    tf2 = add_textbox(slide, Inches(11.4), Inches(7.03), Inches(1.13), Inches(0.3))
    p = tf2.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    add_run(p, f"{page:02d} / {TOTAL_SLIDES:02d}", 8.5, PALETTE["dim"], spc=140)


def _bullet(p, char: str = "\u2022", color: str = PALETTE["accent"], mar_in: float = 0.24) -> None:
    pPr = p._p.get_or_add_pPr()
    mar = int(Inches(mar_in))
    pPr.set("marL", str(mar))
    pPr.set("indent", str(-mar))
    bu_clr = etree.SubElement(pPr, qn("a:buClr"))
    srgb = etree.SubElement(bu_clr, qn("a:srgbClr"))
    srgb.set("val", color.lstrip("#"))
    bu_font = etree.SubElement(pPr, qn("a:buFont"))
    bu_font.set("typeface", "Arial")
    bu_char = etree.SubElement(pPr, qn("a:buChar"))
    bu_char.set("char", char)


def add_bullets(slide, x, y, w, h, items, size=14.5, gap=11, line_spacing=1.24):
    tf = add_textbox(slide, x, y, w, h)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(gap)
        if isinstance(item, tuple):
            strong, detail = item
            add_run(p, strong, size, PALETTE["text"], bold=True)
            if detail:
                add_run(p, "  " + detail, size, PALETTE["muted"])
        else:
            add_run(p, item, size, PALETTE["text_soft"])
        _bullet(p)
    return tf


def add_card(slide, x, y, w, h, title, body, *, num=None, accent=PALETTE["accent"], title_size=15, body_size=11.5):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        shp.adjustments[0] = 0.075
    except Exception:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(PALETTE["panel"])
    shp.line.color.rgb = rgb(PALETTE["line"])
    shp.line.width = Pt(0.75)
    _no_shadow(shp)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.24)
    tf.margin_right = Inches(0.24)
    tf.margin_top = Inches(0.18)
    tf.margin_bottom = Inches(0.16)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p1 = tf.paragraphs[0]
    p1.line_spacing = 1.05
    if num:
        add_run(p1, num + "   ", title_size, accent, bold=True)
    add_run(p1, title, title_size, PALETTE["text"], bold=True)
    p2 = tf.add_paragraph()
    p2.line_spacing = 1.22
    p2.space_before = Pt(6)
    add_run(p2, body, body_size, PALETTE["muted"])
    return shp


def est_chip_width(text: str, pt: float, pad: float = 0.42) -> float:
    return len(text) * pt * 0.0093 + pad  # inches, Arial approximation


def add_chip(slide, x, y, text, size=10.5, color=PALETTE["text_soft"]):
    w = Inches(est_chip_width(text, size))
    h = Inches(0.36)
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        shp.adjustments[0] = 0.5
    except Exception:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb("#2C2A28")
    shp.line.color.rgb = rgb("#47443F")
    shp.line.width = Pt(0.75)
    _no_shadow(shp)
    tf = shp.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.16)
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, text, size, color, bold=True)
    return x + w


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def standard_header(slide, kicker: str, title: str, page: int, size=33):
    add_background(slide)
    add_kicker(slide, kicker)
    next_y = add_title(slide, title, size)
    add_footer(slide, page)
    return next_y


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------
def slide_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s, deep=True)
    add_rect(s, ML, Inches(0.75), Inches(0.16), Inches(0.16), PALETTE["accent"])
    tf = add_textbox(s, ML + Inches(0.32), Inches(0.71), Inches(8.0), Inches(0.3))
    add_run(tf.paragraphs[0], "PITCH  \u00b7  PREPARED FOR VIVEK", 11, PALETTE["accent"], bold=True, spc=220)

    tf = add_textbox(s, ML, Inches(2.0), CW, Inches(1.6))
    add_run(tf.paragraphs[0], "Octopus", 96, PALETTE["text"], bold=True, font=HEAD_FONT, spc=-60)

    tf = add_textbox(s, ML, Inches(3.55), CW, Inches(0.7))
    add_run(tf.paragraphs[0], "Your AI company, in one folder.", 27, PALETTE["text_soft"], bold=True, spc=-10)

    add_rect(s, ML, Inches(4.42), Inches(1.6), Inches(0.06), PALETTE["accent"], rounded=True, radius=0.5)

    tf = add_textbox(s, ML, Inches(4.75), CW, Inches(0.4))
    add_run(tf.paragraphs[0], "A pitch prepared for Vivek  \u00b7  September 2026", 14, PALETTE["muted"])

    x = ML
    for chip in ["Local-first", "Azure OpenAI gpt-6-luna", "Demo Mode \u00b7 no keys", "Runs on a laptop"]:
        x = add_chip(s, x, Inches(5.45), chip) + Inches(0.16)

    tf = add_textbox(s, ML, Inches(6.6), CW, Inches(0.35))
    add_run(tf.paragraphs[0], "github.com/bala2006/octopus", 12, PALETTE["steel"], bold=True)
    add_notes(
        s,
        "Vivek - thanks for the time. In the next few minutes I'll show you Octopus: "
        "instead of one AI coding chat, it's a full company of agents - a CEO, PM, "
        "engineers, a reviewer and QA - working together inside a folder on your laptop. "
        "Everything you'll see is real, local and already built.",
    )


def slide_problem(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    y = standard_header(s, "The problem", "Software is a team sport.\nAI tooling forgot that.", 2, size=31)
    add_bullets(
        s, ML, Inches(y + 0.15), Inches(7.4), Inches(4.0),
        [
            ("One chat window.", "Every AI coding tool is a single agent with no structure - no roles, no channels, no org."),
            ("You are the middleware.", "Planning, review and QA live in your head; you relay every message by hand."),
            ("No accountability.", "Nothing delegates, reports or owns an outcome - the model just answers."),
            ("No control surface.", "What ran, what changed and what it cost: nobody can produce the receipt."),
            ("Cloud by default.", "Your code and your data leave the machine to get any of this."),
        ],
        size=14,
    )
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.55), Inches(2.35), Inches(3.98), Inches(3.5))
    try:
        card.adjustments[0] = 0.06
    except Exception:
        pass
    card.fill.solid()
    card.fill.fore_color.rgb = rgb("#2C2C2C")
    card.line.color.rgb = rgb(PALETTE["accent"])
    card.line.width = Pt(1.0)
    _no_shadow(card)
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.3)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.line_spacing = 1.12
    add_run(p, "The gap isn't intelligence.", 19, PALETTE["text"], bold=True)
    p = tf.add_paragraph()
    p.line_spacing = 1.12
    add_run(p, "It's organization.", 19, PALETTE["accent"], bold=True)
    p = tf.add_paragraph()
    p.space_before = Pt(14)
    p.line_spacing = 1.3
    add_run(p, "Companies solved coordination 100 years ago: roles, channels, review, accountability. "
               "Software teams never got it.", 12, PALETTE["muted"])
    add_notes(
        s,
        "The market is crowded with single-agent copilots. They're smart but unstructured: "
        "the human ends up being the project manager, the reviewer and the QA lead. "
        "And there's no audit trail or cost visibility. We think the missing layer is organization, "
        "not more intelligence.",
    )


def slide_idea(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    y = standard_header(s, "The idea", "A small company of AI agents.", 3)
    x = ML
    for role in ["CEO", "PM", "Engineer", "Reviewer", "QA"]:
        x = add_chip(s, x, Inches(y + 0.05), role) + Inches(0.14)
    add_bullets(
        s, ML, Inches(y + 0.7), Inches(7.4), Inches(3.6),
        [
            ("Pick a folder, pick a team, give it a goal.", "14 templates wire the org for you - or design one on the canvas."),
            ("They do the work.", "Plan, debate and delegate along typed channels; write real files; review each other; run the tests."),
            ("You stay the boss.", "Watch it live, approve risky steps, interject, and continue the run like a chat afterwards."),
        ],
        size=14.5,
    )
    add_card(
        s, Inches(8.55), Inches(y + 0.05), Inches(3.98), Inches(3.4),
        "Local-first by design",
        "Everything Octopus knows lives in <folder>/.octopus/ - a SQLite DB, plans, exports "
        "and a .gitignore of *. Like .git, re-opening the folder brings the whole company back. "
        "Nothing is uploaded; Demo Mode needs no keys.",
        title_size=15,
    )
    add_notes(
        s,
        "You pick a folder - that's the sandbox - pick one of 14 team templates, and type a goal like "
        "'build a todo app with login'. The agents plan, argue, delegate and code. Files land in your real "
        "folder, and .octopus/ keeps all state locally, the way .git does.",
    )


def slide_how(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    y = standard_header(s, "How a run works", "Goal in. Software out.", 4)
    steps = [
        ("01", "Goal", "Delivered to the entry agent's mailbox"),
        ("02", "Plan & debate", "Proposal - objection - explicit consensus"),
        ("03", "Build", "write_file / run_code inside the sandbox"),
        ("04", "Review", "Verdicts, revisions, QA runs real tests"),
        ("05", "Ship", "Real files + a run report you can replay"),
    ]
    sw, gap = Inches(2.19), Inches(0.19)
    x = ML
    top = Inches(y + 0.35)
    for i, (num, title, body) in enumerate(steps):
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, sw, Inches(2.15))
        try:
            shp.adjustments[0] = 0.09
        except Exception:
            pass
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(PALETTE["panel"])
        shp.line.color.rgb = rgb(PALETTE["line"])
        shp.line.width = Pt(0.75)
        _no_shadow(shp)
        tf = shp.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.18)
        tf.margin_right = Inches(0.18)
        tf.margin_top = Inches(0.16)
        p = tf.paragraphs[0]
        add_run(p, num, 12, PALETTE["accent"], bold=True, spc=120)
        p = tf.add_paragraph()
        p.space_before = Pt(6)
        add_run(p, title, 15, PALETTE["text"], bold=True)
        p = tf.add_paragraph()
        p.space_before = Pt(5)
        p.line_spacing = 1.22
        add_run(p, body, 10.5, PALETTE["muted"])
        if i < len(steps) - 1:
            arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + sw + Inches(0.015), top + Inches(0.93),
                                       Inches(0.16), Inches(0.16))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = rgb(PALETTE["accent"])
            arrow.line.fill.background()
            _no_shadow(arrow)
        x += sw + gap
    tf = add_textbox(s, ML, Inches(y + 2.85), CW, Inches(0.9))
    p = tf.paragraphs[0]
    p.line_spacing = 1.35
    add_run(p, "One runtime per run.  ", 13, PALETTE["text"], bold=True)
    add_run(p, "Agents reply with a validated JSON action schema; the server only delivers messages along "
               "channels it knows about, every action passes the permission gate, and every event is persisted "
               "with a sequence number so the whole run can be replayed frame by frame.", 13, PALETTE["muted"])
    x = ML
    for chip in ["Explicit consensus", "Review verdicts", "Task board", "Loop detector"]:
        x = add_chip(s, x, Inches(y + 3.75), chip) + Inches(0.14)
    add_notes(
        s,
        "This is the core loop. A goal enters the entry agent's mailbox; the team debates a plan - consensus "
        "has to be explicit, both sides must agree - then engineers write files in the sandbox, an architect "
        "reviews with itemized comments, and QA actually runs the tests. Everything is a validated action, "
        "enforced server-side, and recorded so you can scrub back through the timeline.",
    )


def slide_product(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    y = standard_header(s, "Product", "A whole product, not a demo.", 5)
    cards = [
        ("Org canvas", "Departments, managers and typed channels - delegate, report, review, debate, consult. Undo/redo, auto-layout, JSON import/export, autosave."),
        ("Live run", "Animated graph of who's thinking, messages moving along edges, message feed, task board, tool log, approvals, pause / step / interject, replay timeline."),
        ("Runs are chats", "When a run finishes, type 'add a dark mode toggle' in the same box - the same run re-opens with its history, task board and files."),
        ("Artifacts & preview", "Every file version with a diff, author and reason, plus revert and ZIP. Live preview of HTML, Markdown, images and PDFs."),
    ]
    cw_, ch = Inches(5.77), Inches(1.92)
    gx, gy = Inches(0.35), Inches(0.3)
    positions = [(ML, Inches(y + 0.2)), (ML + cw_ + gx, Inches(y + 0.2)),
                 (ML, Inches(y + 0.2) + ch + gy), (ML + cw_ + gx, Inches(y + 0.2) + ch + gy)]
    for (title, body), (px, py) in zip(cards, positions):
        add_card(s, px, py, cw_, ch, title, body)
    add_notes(
        s,
        "Four surfaces make this a product: the org canvas where you design the company, the live run view "
        "where you watch and intervene, runs-as-chats so work accumulates instead of restarting, and artifacts "
        "where every change is versioned and previewable.",
    )


def slide_control(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    y = standard_header(s, "Control & safety", "Autonomy with a leash.", 6)
    rows, cols = 5, 5
    tbl_x, tbl_y = ML, Inches(y + 0.2)
    tbl_w, tbl_h = Inches(7.5), Inches(2.55)
    table = s.shapes.add_table(rows, cols, tbl_x, tbl_y, tbl_w, tbl_h).table
    table.first_row = False
    table.horz_banding = False
    table.columns[0].width = Inches(2.7)
    for c in range(1, 5):
        table.columns[c].width = Inches(1.2)
    table.rows[0].height = Inches(0.5)
    for r in range(1, 5):
        table.rows[r].height = Inches(0.51)
    header = ["Action", "read_only", "plan", "ask", "danger"]
    data = [
        ("read / list files", ["yes", "shadow first", "yes", "yes + secrets"]),
        ("write_file", ["no", "shadow", "approve", "yes"]),
        ("run_code", ["no", "no", "approve", "any command"]),
        ("mcp_call", ["no", "no", "approve", "yes"]),
    ]

    def style_cell(cell, text, *, bold, color, fill, size=11.5, align=PP_ALIGN.CENTER):
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
        cell.margin_left = Inches(0.1)
        cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.05)
        cell.margin_bottom = Inches(0.05)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        add_run(p, text, size, color, bold=bold)

    for c, htext in enumerate(header):
        style_cell(table.cell(0, c), htext, bold=True, color=PALETTE["text"],
                   fill="#33302D", size=12,
                   align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
    allowed = {"yes": PALETTE["olive"], "yes + secrets": PALETTE["olive"], "shadow": PALETTE["steel"],
               "shadow first": PALETTE["steel"], "approve": PALETTE["accent"], "any command": PALETTE["olive"]}
    for r, (action, cells) in enumerate(data, start=1):
        style_cell(table.cell(r, 0), action, bold=True, color=PALETTE["text_soft"],
                   fill="#2C2A28", size=11.5, align=PP_ALIGN.LEFT)
        for c, val in enumerate(cells, start=1):
            col = allowed.get(val, "#7A766F")
            style_cell(table.cell(r, c), val, bold=(val in ("yes", "approve")), color=col,
                       fill="#262421", size=10.5)
    add_line_of_text(s, tbl_x, tbl_y + tbl_h + Inches(0.12), tbl_w, Inches(0.3),
                     "Effective level = min(run level, agent override).", 10.5, PALETTE["dim"], italic=True)

    add_bullets(
        s, Inches(8.55), Inches(y + 0.2), Inches(3.98), Inches(4.2),
        [
            ("Approvals with a diff.", "Writes, commands and MCP calls block for sign-off; 'Always allow' scopes to that agent and action."),
            ("Budgets everywhere.", "Turns, tokens, cost and wall-clock - plus a loop detector that auto-pauses runaway runs."),
            ("Sandboxed execution.", "No shell, scrubbed environment, rlimits, no network, allowlist - or Docker isolation."),
            ("Your secrets stay put.", "Keys are Fernet-encrypted on your machine and never returned to the browser."),
        ],
        size=12,
        gap=10,
    )
    add_notes(
        s,
        "Autonomy is a dial, not a switch. read_only can't touch anything, plan writes to a shadow you apply "
        "later, ask blocks on your approval with a full diff, danger is unrestricted. Budgets and the loop "
        "detector are backstops, and commands run with no shell, rlimits and no network by default.",
    )


def slide_model(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    y = standard_header(s, "The model layer", "One model, done properly.", 7)
    add_bullets(
        s, ML, Inches(y + 0.2), Inches(7.4), Inches(4.0),
        [
            ("Azure OpenAI gpt-6-luna.", "Called straight over the v1 Responses API - any endpoint shape works, no api-version quirks."),
            ("Exact usage, exact cost.", "Uncached input, cached input, cache writes and output (incl. reasoning), priced at real rates."),
            ("Reasoning effort per agent.", "none / low / medium / high / x-high / max - tune how hard each role thinks."),
            ("Retries that behave.", "429s and 5xxs back off with jitter - but only before the first token streams."),
            ("Demo Mode.", "A scripted, offline company that reacts to the real inbox. No keys, no network, free forever."),
        ],
        size=14,
    )
    add_card(
        s, Inches(8.55), Inches(y + 0.2), Inches(3.98), Inches(3.6),
        "Priced in your currency",
        "Per 1M tokens at gpt-6-luna Standard rates:\n"
        "$0.10 input  \u00b7  $0.01 cached\n"
        "$0.125 cache writes  \u00b7  $0.50 output\n\n"
        "Long prompts over 272K bill at 2x input. Totals are stored in USD and shown in INR, EUR, GBP ... "
        "using live ECB rates, cached for 6 hours and offline-safe.",
        title_size=15,
    )
    add_notes(
        s,
        "One deployment, wired properly: the v1 Responses API, reasoning effort as a first-class control, and "
        "usage read straight from Azure rather than estimated - so the cost meter is exact, down to cache writes. "
        "You can see spend in rupees, dollars or any currency. And if you don't connect a key at all, Demo Mode "
        "still runs a full company offline.",
    )


def slide_stack(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    y = standard_header(s, "Engineering", "Built like real software.", 8)
    cards = [
        ("Backend", "FastAPI \u00b7 SQLAlchemy + Alembic \u00b7 async SQLite \u00b7 WebSockets with snapshot + replay \u00b7 structlog \u00b7 MCP client \u00b7 rate limiting and strict CORS."),
        ("Frontend", "React 18 + TypeScript \u00b7 Vite \u00b7 Tailwind + shadcn-style UI \u00b7 React Flow canvas \u00b7 Zustand with history \u00b7 Monaco \u00b7 framer-motion. Fonts bundled - zero CDNs."),
        ("Typed by contract", "The TypeScript client is generated from FastAPI's OpenAPI schema in CI, so backend and UI can't drift apart."),
        ("Tested & shipped", "pytest (68 tests) + vitest + Playwright E2E on every push, then CD publishes ghcr.io images. Docker compose for a fully containered run."),
    ]
    cw_, ch = Inches(5.77), Inches(1.92)
    gx, gy = Inches(0.35), Inches(0.3)
    positions = [(ML, Inches(y + 0.2)), (ML + cw_ + gx, Inches(y + 0.2)),
                 (ML, Inches(y + 0.2) + ch + gy), (ML + cw_ + gx, Inches(y + 0.2) + ch + gy)]
    for (title, body), (px, py) in zip(cards, positions):
        add_card(s, px, py, cw_, ch, title, body)
    add_notes(
        s,
        "This isn't a prototype. 23,000+ lines across backend and frontend, migrations, a generated typed API "
        "client, unit tests, end-to-end Playwright tests that continue a finished run, CI on every push and CD "
        "publishing container images. It also runs with no CDN at all, so it works fully offline.",
    )


def slide_proof(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    y = standard_header(s, "Proof", "Proof, not promises.", 9)
    add_bullets(
        s, ML, Inches(y + 0.2), Inches(7.4), Inches(4.0),
        [
            ("1-minute demo, real footage.", "Recorded from the actual app in Demo Mode by scripts/demo-video/record.mjs - regenerate with one command."),
            ("14 team templates.", "Software startup, web studio, game studio, SaaS launch, security audit, research lab, support desk ..."),
            ("End-to-end tests that matter.", "E2E covers the happy path, run continuation, approvals, org generation and live hiring."),
            ("One command to run.", "make setup && make start -> http://localhost:8000. Python 3.11 + Node 20, nothing else."),
            ("Docs that keep up.", "README, ARCHITECTURE, DEMO script and an illustrated in-app Guide covering every feature."),
        ],
        size=14,
    )
    stats = [("216", "files in repo"), ("23k+", "lines of code"), ("14", "team templates"), ("68", "pytest tests")]
    sx, sy = Inches(8.55), Inches(y + 0.2)
    bw, bh, bgap = Inches(1.94), Inches(1.6), Inches(0.22)
    for i, (big, small) in enumerate(stats):
        px = sx + (i % 2) * (bw + bgap)
        py = sy + (i // 2) * (bh + bgap)
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, py, bw, bh)
        try:
            shp.adjustments[0] = 0.1
        except Exception:
            pass
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(PALETTE["panel"])
        shp.line.color.rgb = rgb(PALETTE["line"])
        shp.line.width = Pt(0.75)
        _no_shadow(shp)
        tf = shp.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.14)
        tf.margin_right = Inches(0.12)
        tf.margin_top = Inches(0.24)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_run(p, big, 30, PALETTE["accent"], bold=True)
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(4)
        add_run(p, small, 10.5, PALETTE["muted"])
    add_notes(
        s,
        "Everything I've shown is in the repo today: the demo video is recorded from the real app, the test suite "
        "runs on every push, and a new contributor is one make command away from a running company.",
    )


def slide_vision(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    y = standard_header(s, "Where this goes", "Every team gets an AI company.", 10)
    add_bullets(
        s, ML, Inches(y + 0.2), Inches(7.6), Inches(4.2),
        [
            ("Beyond code.", "The same org model runs content studios, research labs, support desks - the templates already exist."),
            ("A marketplace of teams.", "Publish, fork and rate companies the way people share Dockerfiles today."),
            ("Hosted control plane.", "Same local-first core, with team sharing, SSO and org analytics for companies."),
            ("The folder is the contract.", "Portable, auditable, diff-able - your agents and your output never lock into our cloud."),
        ],
        size=14.5,
    )
    add_card(
        s, Inches(8.75), Inches(y + 0.2), Inches(3.78), Inches(3.4),
        "Why now",
        "Models finally handle multi-step tool use reliably, MCP made tools plug-and-play, and every team "
        "already wants leverage. The winning product won't be a better chat box - it'll be the thing that "
        "organizes the work.",
        title_size=15,
    )
    add_notes(
        s,
        "Octopus is a platform, not a coding toy. The org abstraction already generalizes - we ship templates "
        "for content, research and support. The roadmap is a team marketplace plus a hosted control plane, "
        "while keeping the folder as the portable source of truth.",
    )


def slide_ask(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    y = standard_header(s, "The ask", "Vivek, here is the ask.", 11)
    asks = [
        ("01", "Be our first power user", "Run one real goal end to end this week - a real project, not a toy - and tear the experience apart."),
        ("02", "Open doors", "Intros to 3 teams who would run an AI company daily: agencies, dev shops, studios, research groups."),
        ("03", "Advise us", "Pricing and packaging, open-source go-to-market, and the path to a hosted version without betraying local-first."),
    ]
    cw_, ch = Inches(3.77), Inches(2.6)
    gap = Inches(0.31)
    x = ML
    top = Inches(y + 0.25)
    for num, title, body in asks:
        add_card(s, x, top, cw_, ch, title, body, num=num, title_size=15.5, body_size=12)
        x += cw_ + gap
    add_line_of_text(
        s, ML, top + ch + Inches(0.28), CW, Inches(0.5),
        "In return: the 90-day roadmap (hosted teams, marketplace, org analytics) and a build partner "
        "who ships weekly and shares every number.", 13.5, PALETTE["muted"], line_spacing=1.3,
    )
    add_notes(
        s,
        "Three asks: be the first power user and give me brutal product feedback; make a handful of intros to "
        "teams who'd use this daily; and advise us on pricing and GTM. In return you get the roadmap, the "
        "metrics, and a team that ships every week.",
    )


def slide_closing(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s, deep=True)
    tf = add_textbox(s, Inches(1.2), Inches(2.35), Inches(10.93), Inches(2.2))
    for i, line in enumerate(["Assemble an AI company.", "Give it a goal. Ship."]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.15
        p.alignment = PP_ALIGN.CENTER
        add_run(p, line, 44, PALETTE["text"], bold=True, spc=-20)
    add_rect(s, Inches(6.07), Inches(4.55), Inches(1.2), Inches(0.06), PALETTE["accent"], rounded=True, radius=0.5)
    tf = add_textbox(s, Inches(1.2), Inches(4.9), Inches(10.93), Inches(0.5))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "github.com/bala2006/octopus", 16, PALETTE["steel"], bold=True)
    tf = add_textbox(s, Inches(1.2), Inches(5.5), Inches(10.93), Inches(0.5))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "Thank you, Vivek.", 15, PALETTE["muted"])
    add_notes(
        s,
        "That's Octopus: an AI company in one folder - structured, supervised and local. "
        "I'd love to get you into it this week. What's the best next step for you?",
    )


def build(output: Path) -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_title(prs)
    slide_problem(prs)
    slide_idea(prs)
    slide_how(prs)
    slide_product(prs)
    slide_control(prs)
    slide_model(prs)
    slide_stack(prs)
    slide_proof(prs)
    slide_vision(prs)
    slide_ask(prs)
    slide_closing(prs)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return output


def main() -> None:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "octopus-pitch-for-vivek.pptx"
    path = build(out)
    print(f"wrote {path} ({path.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
